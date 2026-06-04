import os
import re
import docx
from docx.shared import RGBColor
from docx.enum.style import WD_STYLE_TYPE

def is_code_line(line: str) -> bool:
    """Heuristic function to determine if a line of text is a Python code snippet."""
    line = line.strip()
    if not line: return False
    
    # Text headers or explanations usually have " : "
    if " : " in line and not line.startswith("#"):
        return False
        
    code_prefixes = (
        'import ', 'from ', 'def ', 'class ', 'if ', 'elif ', 'else:', 
        'try:', 'except', 'finally:', 'for ', 'while ', 'return ', 'yield ', 
        '#', '>>>', '"""', "'''", 'del ', 'pass', 'break', 'continue',
        '@', 'print(', 'cd ', 'python ', 'super().', 'self.', 'my_', 
        'list(', 'set(', 'dict(', 'tuple('
    )
    if line.startswith(code_prefixes): return True
    
    # Variable assignments (e.g., x = 5 or name = 'John')
    if re.search(r'^[a-zA-Z_][a-zA-Z0-9_]*\s*[\+\-\*\/\%]?=\s*', line): return True
    
    # Method calls (e.g., my_list.append())
    if re.search(r'^[a-zA-Z_][a-zA-Z0-9_]*\.[a-zA-Z_]', line): return True
    
    # Dict/Set items (e.g., 'key': value)
    if re.search(r'^[\'"][a-zA-Z0-9_ -]+[\'"]\s*:', line): return True
    
    # Lone brackets
    if line in ('{', '}', '[', ']', '(', ')'): return True
    
    # Variable evaluations (e.g., typing 'name' in a REPL)
    if re.match(r'^[a-zA-Z_][a-zA-Z0-9_]*$', line) and line.islower(): return True
    
    # Mathematical expressions
    if re.search(r'^[a-zA-Z_][a-zA-Z0-9_]*\s+[\&\|\^\-]\s+[a-zA-Z_]', line): return True
    
    # Tracebacks & Errors
    if line.startswith('Traceback ') or line.startswith('File "<') or ('Error:' in line and not line.startswith('The ')): return True
    
    # Strings checking "in" or equalities
    if line.startswith(("'", '"')) and (" in " in line or "==" in line): return True
    
    # Slices & Lists/Tuples outputs (e.g., [1, 2, 3] or ('Laptop', 990))
    if re.match(r'^[a-zA-Z_]+\[.*\]$', line) or re.match(r'^\[.*\]$', line) or re.match(r'^\(.*\)$', line): return True

    return False

input_path = r"Summary Notes.docx"
output_path = r"Formatted_Summary_Notes.docx"

def format_document_in_place(input_doc_path: str, output_doc_path: str) -> None:
    """Reads a Word document, identifies lines that look like code snippets, and 
    applies a custom style to them while preserving the original formatting of non-code lines."""
    if not os.path.exists(input_doc_path):
        print(f"Error: Could not find '{input_doc_path}'.")
        return

    print(f"Opening {input_doc_path} to edit in-place...")
    doc = docx.Document(input_doc_path)
    
    # Set up the custom "Code Emphasis" style inside the existing document
    styles = doc.styles
    try:
        code_style = styles.add_style('Code Emphasis', WD_STYLE_TYPE.CHARACTER)
        code_style.font.italic = True
        code_style.font.color.rgb = RGBColor(0x4F, 0x81, 0xBD) 
        code_style.font.name = 'Courier New'
    except Exception:
        # If the style was already created on a previous run, just skip creating it
        pass

    print("Scanning and formatting code lines...")
    for para in doc.paragraphs:
        original_text = para.text
        
        # If the heuristic flags it as code, we overwrite that specific paragraph's formatting
        if is_code_line(original_text):
            para.clear() # Removes the old plain-text formatting
            new_run = para.add_run(original_text.strip())
            new_run.style = 'Code Emphasis'
        
        # If it is NOT code, we do absolutely nothing. The original formatting is preserved!

    doc.save(output_doc_path)
    print(f"Success! Document saved to: {os.path.abspath(output_doc_path)}")

if __name__ == '__main__':
    format_document_in_place(input_path, output_path)